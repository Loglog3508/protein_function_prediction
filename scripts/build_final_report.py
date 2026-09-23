"""Build the final competition report as a formatted DOCX."""

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "reports/final_report.md"
DESTINATION = ROOT / "reports/蛋白质功能预测分析报告.docx"
HEADER_TITLE = "蛋白质功能预测分析报告"


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, code, separate, end])


def configure_font(style, *, chinese: str, western: str, size: float, bold=False):
    style.font.name = western
    style.font.size = Pt(size)
    style.font.bold = bold
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), chinese)


def set_run_east_asia(run, font_name: str) -> None:
    properties = run._element.get_or_add_rPr()
    fonts = properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.insert(0, fonts)
    fonts.set(qn("w:eastAsia"), font_name)


def configure_page(section) -> None:
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3.7)
    section.bottom_margin = Cm(4.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(2.7)
    section.footer_distance = Cm(3.0)


def set_page_numbering(section, *, number_format: str, start: int) -> None:
    properties = section._sectPr
    existing = properties.find(qn("w:pgNumType"))
    if existing is not None:
        properties.remove(existing)
    page_number = OxmlElement("w:pgNumType")
    page_number.set(qn("w:fmt"), number_format)
    page_number.set(qn("w:start"), str(start))
    columns = properties.find(qn("w:cols"))
    properties.insert(
        properties.index(columns) if columns is not None else len(properties),
        page_number,
    )


def configure_header_footer(section, *, number_format: str, start: int) -> None:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    header = section.header.paragraphs[0]
    header.text = ""
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run(HEADER_TITLE)
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)
    set_run_east_asia(run, "宋体")

    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(footer, "PAGE")
    set_page_numbering(section, number_format=number_format, start=start)


def configure_document(document: Document) -> None:
    configure_page(document.sections[0])
    settings = document.settings._element
    zoom = settings.find(qn("w:zoom"))
    if zoom is not None:
        zoom.set(qn("w:percent"), "100")
    language = settings.find(qn("w:themeFontLang"))
    if language is not None:
        language.set(qn("w:eastAsia"), "zh-CN")

    normal = document.styles["Normal"]
    configure_font(normal, chinese="宋体", western="Times New Roman", size=12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(20)
    normal.paragraph_format.first_line_indent = Cm(0.85)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    heading_specs = {
        "Heading 1": (16, 30, 36),
        "Heading 2": (14, 18, 24),
        "Heading 3": (12, 12, 15),
        "Heading 4": (12, 6, 9),
    }
    for name, (size, before, after) in heading_specs.items():
        style = document.styles[name]
        configure_font(style, chinese="黑体", western="Arial", size=size, bold=True)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        style.paragraph_format.line_spacing = Pt(20)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    front_title = document.styles.add_style("Front Title", WD_STYLE_TYPE.PARAGRAPH)
    configure_font(front_title, chinese="黑体", western="Arial", size=16, bold=True)
    front_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    front_title.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    front_title.paragraph_format.line_spacing = Pt(20)
    front_title.paragraph_format.space_after = Pt(36)

    figure_caption = document.styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    configure_font(figure_caption, chinese="黑体", western="Times New Roman", size=10.5)
    figure_caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure_caption.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    figure_caption.paragraph_format.line_spacing = Pt(20)
    figure_caption.paragraph_format.space_after = Pt(10)

    table_caption = document.styles.add_style("Table Caption", WD_STYLE_TYPE.PARAGRAPH)
    configure_font(table_caption, chinese="黑体", western="Times New Roman", size=10.5)
    table_caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table_caption.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    table_caption.paragraph_format.line_spacing = Pt(20)
    table_caption.paragraph_format.keep_with_next = True

    reference = document.styles.add_style("Reference", WD_STYLE_TYPE.PARAGRAPH)
    configure_font(reference, chinese="宋体", western="Times New Roman", size=10.5)
    reference.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    reference.paragraph_format.left_indent = Cm(0.74)
    reference.paragraph_format.first_line_indent = Cm(-0.74)
    reference.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    reference.paragraph_format.line_spacing = Pt(20)
    reference.paragraph_format.space_before = Pt(0)
    reference.paragraph_format.space_after = Pt(0)


def remove_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "nil")
        borders.append(border)
    insert_table_borders(properties, borders)


def insert_table_borders(properties, borders) -> None:
    for successor in ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook"):
        element = properties.find(qn(successor))
        if element is not None:
            properties.insert(properties.index(element), borders)
            return
    properties.append(borders)


def add_cover(document: Document, title: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(50)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    for text, size in (
        ("2026 年全国大学生“海豚杯”", 22),
        ("数智应用创新大赛", 22),
        ("数智分析赛报告", 22),
    ):
        run = paragraph.add_run(text)
        run.add_break()
        run.bold = True
        run.font.size = Pt(size)
        set_run_east_asia(run, "黑体")

    label = document.add_paragraph("题目")
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label.paragraph_format.space_before = Pt(25)
    label.runs[0].bold = True
    label.runs[0].font.size = Pt(16)
    set_run_east_asia(label.runs[0], "黑体")

    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_paragraph.paragraph_format.space_before = Pt(14)
    title_paragraph.paragraph_format.space_after = Pt(36)
    title_paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = title_paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(18)
    set_run_east_asia(run, "黑体")

    table = document.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    remove_table_borders(table)
    for row, label_text in zip(
        table.rows,
        ("团队名称：", "参赛高校：", "参赛队员：", "指导教师："),
        strict=True,
    ):
        row.cells[0].width = Cm(3)
        row.cells[1].width = Cm(8)
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        label_run = row.cells[0].paragraphs[0].add_run(label_text)
        value_run = row.cells[1].paragraphs[0].add_run("________________________")
        for run in (label_run, value_run):
            run.font.size = Pt(14)
            set_run_east_asia(run, "宋体")

    date = document.add_paragraph("2026 年  9 月")
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.paragraph_format.space_before = Pt(42)
    date.runs[0].font.size = Pt(14)
    set_run_east_asia(date.runs[0], "宋体")


def add_toc(document: Document) -> None:
    heading = document.add_paragraph("目  录", style="Front Title")
    heading.paragraph_format.space_before = Pt(0)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0)
    add_field(paragraph, 'TOC \\o "1-3" \\h \\z \\u')


def add_inline_text(paragraph, text: str, *, citations: bool = True) -> None:
    pattern = r"(`[^`]+`|\*\*[^*]+\*\*|\[(?:\d+(?:[-,]\d+)*)\])"
    for part in re.split(pattern, text):
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10.5)
        elif part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif citations and re.fullmatch(r"\[(?:\d+(?:[-,]\d+)*)\]", part):
            run = paragraph.add_run(part)
            run.font.superscript = True
        else:
            paragraph.add_run(part)


def set_three_line_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge, value, size in (
        ("top", "single", "12"),
        ("left", "nil", "0"),
        ("bottom", "single", "12"),
        ("right", "nil", "0"),
        ("insideH", "nil", "0"),
        ("insideV", "nil", "0"),
    ):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), value)
        border.set(qn("w:sz"), size)
        borders.append(border)
    insert_table_borders(properties, borders)

    for cell in table.rows[0].cells:
        cell_properties = cell._tc.get_or_add_tcPr()
        cell_borders = OxmlElement("w:tcBorders")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        cell_borders.append(bottom)
        cell_properties.append(cell_borders)


def parse_table(lines: list[str], start: int, document: Document) -> int:
    rows = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            rows.append(cells)
        index += 1
    if not rows:
        return index

    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_three_line_borders(table)
    if rows[0][0] == "模型":
        column_widths = [6.5, 2.0, 4.5, 3.0]
    elif rows[0][0] == "阈值策略":
        column_widths = [4.5, 4.0, 4.0, 3.5]
    else:
        column_widths = [16.0 / len(rows[0])] * len(rows[0])
    for column, width in zip(table.columns, column_widths, strict=True):
        column.width = Cm(width)
    for row_index, values in enumerate(rows):
        row_properties = table.rows[row_index]._tr.get_or_add_trPr()
        row_properties.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            row_properties.append(OxmlElement("w:tblHeader"))
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.width = Cm(column_widths[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            add_inline_text(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(10.5)
                set_run_east_asia(run, "宋体")
                if row_index == 0:
                    run.bold = True
    document.add_paragraph()
    return index


def add_front_matter(document: Document, lines: list[str]) -> int:
    summary_index = next(
        index for index, line in enumerate(lines) if line.strip() == "## 分析结果概要"
    )
    body_index = next(
        index
        for index in range(summary_index + 1, len(lines))
        if lines[index].startswith("## ")
    )

    section = document.add_section(WD_SECTION.NEW_PAGE)
    configure_page(section)
    configure_header_footer(section, number_format="upperRoman", start=1)
    heading = document.add_paragraph("分析结果概要", style="Heading 1")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for raw in lines[summary_index + 1 : body_index]:
        stripped = raw.strip()
        if stripped:
            paragraph = document.add_paragraph()
            add_inline_text(paragraph, stripped)
    document.add_page_break()
    add_toc(document)
    return body_index


def add_body(document: Document, lines: list[str], start: int) -> None:
    section = document.add_section(WD_SECTION.NEW_PAGE)
    configure_page(section)
    configure_header_footer(section, number_format="decimal", start=1)

    index = start
    chapter_number = 0
    figure_number = 0
    first_chapter = True
    in_references = False
    while index < len(lines):
        stripped = lines[index].strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith("|"):
            index = parse_table(lines, index, document)
            continue
        if stripped.startswith("## "):
            if not first_chapter:
                document.add_page_break()
            first_chapter = False
            heading_text = stripped[3:].strip()
            in_references = heading_text == "参考文献"
            if not in_references:
                chapter_number += 1
                figure_number = 0
            heading = document.add_heading(heading_text, level=1)
            if in_references:
                heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif stripped.startswith("### "):
            document.add_heading(stripped[4:].strip(), level=2)
        elif stripped.startswith("#### "):
            document.add_heading(stripped[5:].strip(), level=3)
        elif re.fullmatch(r"表\d+\.\d+\s+.+", stripped):
            paragraph = document.add_paragraph(style="Table Caption")
            add_inline_text(paragraph, stripped)
        elif stripped.startswith("!["):
            match = re.fullmatch(r"!\[(.+)]\((.+)\)", stripped)
            if match:
                figure_number += 1
                image_path = ROOT / "reports" / match.group(2)
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.first_line_indent = Cm(0)
                paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                paragraph.paragraph_format.keep_with_next = True
                paragraph.add_run().add_picture(str(image_path), width=Cm(15.5))
                caption = document.add_paragraph(style="Figure Caption")
                caption.add_run(f"图{chapter_number}.{figure_number}  {match.group(1)}")
        elif in_references and re.match(r"^\[\d+]", stripped):
            paragraph = document.add_paragraph(style="Reference")
            add_inline_text(paragraph, stripped, citations=False)
        else:
            paragraph = document.add_paragraph()
            add_inline_text(paragraph, stripped)
        index += 1


def build_report() -> Path:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    title = lines[0].removeprefix("# ").strip()
    document = Document()
    configure_document(document)
    document.core_properties.title = title
    document.core_properties.subject = "蛋白质功能预测多标签分类竞赛分析报告"
    document.core_properties.author = "参赛团队"
    document.core_properties.keywords = "蛋白质功能预测, 多标签分类, k-mer, TF-IDF"
    add_cover(document, title)
    body_index = add_front_matter(document, lines)
    add_body(document, lines, body_index)
    DESTINATION.parent.mkdir(parents=True, exist_ok=True)
    document.save(DESTINATION)
    return DESTINATION


if __name__ == "__main__":
    print(build_report())
